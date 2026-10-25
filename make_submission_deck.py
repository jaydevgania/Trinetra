"""
Trinetra Chaos Console — 6-Slide Submission Deck
Output: D:\VJTI\Trinetra_Submission.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Palette ───────────────────────────────────────────────────────────────────
BG       = RGBColor(0x09, 0x0D, 0x14)
BG2      = RGBColor(0x11, 0x18, 0x24)
CARD     = RGBColor(0x16, 0x1E, 0x2E)
CARD2    = RGBColor(0x1C, 0x26, 0x38)
CYAN     = RGBColor(0x00, 0xD4, 0xFF)
GOLD     = RGBColor(0xFF, 0xC2, 0x00)
GREEN    = RGBColor(0x00, 0xE0, 0x70)
RED      = RGBColor(0xFF, 0x45, 0x45)
PURPLE   = RGBColor(0xA0, 0x6C, 0xFF)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
SILVER   = RGBColor(0xB0, 0xBE, 0xCC)
DIM      = RGBColor(0x55, 0x66, 0x77)

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]

# ── Primitives ────────────────────────────────────────────────────────────────

def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s

def rect(s, l, t, w, h, fill=None, line_color=None, line_w=Pt(0)):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    sh = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    else:
        sh.fill.background()
    if line_color:
        sh.line.color.rgb = line_color
        sh.line.width = line_w
    else:
        sh.line.fill.background()
    return sh

def label(s, text, l, t, w, h,
          size=14, bold=False, italic=False,
          color=WHITE, align=PP_ALIGN.LEFT, wrap=True):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tb.word_wrap = wrap
    tf = tb.text_frame
    tf.word_wrap = wrap
    p  = tf.paragraphs[0]
    p.alignment = align
    r  = p.add_run()
    r.text = text
    r.font.size   = Pt(size)
    r.font.bold   = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return tb

def multiline(s, lines, l, t, w, h, size=13, color=WHITE, gap=4, bold=False):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tb.word_wrap = True
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for text in lines:
        if first:
            p = tf.paragraphs[0]; first = False
        else:
            p = tf.add_paragraph()
        p.space_before = Pt(gap)
        r = p.add_run()
        r.text = text
        r.font.size  = Pt(size)
        r.font.bold  = bold
        r.font.color.rgb = color
    return tb

def bar(s, l, t, w, color=CYAN, h=0.045):
    rect(s, l, t, w, float(h), fill=color)

def pill(s, text, l, t, w, h, bg, fg=BG, size=11, bold=True):
    rect(s, l, t, w, h, fill=bg)
    label(s, text, l, t+0.02, w, h-0.04,
          size=size, bold=bold, color=fg, align=PP_ALIGN.CENTER)

def kpi_box(s, big, unit, sub, l, t, w, h, accent):
    rect(s, l, t, w, h, fill=CARD)
    bar(s, l, t, w, color=accent, h=0.055)
    label(s, big,  l, t+0.1,  w, 0.75, size=46, bold=True,
          color=accent, align=PP_ALIGN.CENTER)
    label(s, unit, l, t+0.8,  w, 0.35, size=14,
          color=WHITE, align=PP_ALIGN.CENTER)
    label(s, sub,  l, t+1.12, w, 0.3,  size=10,
          color=SILVER, align=PP_ALIGN.CENTER)

def side_card(s, icon, title, body, l, t, w=5.8, h=1.45,
              accent=CYAN):
    rect(s, l, t, w, h, fill=CARD2)
    bar(s, l, t, 0.06, color=accent, h=h)
    label(s, icon,  l+0.22, t+0.12,  0.6,  0.5, size=22)
    label(s, title, l+0.85, t+0.12,  w-1.0, 0.38,
          size=13, bold=True, color=accent)
    label(s, body,  l+0.85, t+0.5,  w-1.0, h-0.6,
          size=11.5, color=SILVER)

def section_header(s, tag, title, accent=CYAN):
    bar(s, 0, 0, 13.33, color=accent, h=0.07)
    rect(s, 0, 7.18, 13.33, 0.32, fill=CARD)
    label(s, "TRINETRA CHAOS CONSOLE  ·  Innovise 2026  ·  Blockchain PS3",
          0.35, 7.21, 9, 0.26, size=8, color=DIM)
    label(s, tag, 0.38, 0.18, 3, 0.38,
          size=10, bold=True, color=accent)
    label(s, title, 0.38, 0.5, 12, 0.7,
          size=28, bold=True, color=WHITE)
    bar(s, 0.38, 1.22, 5.5, accent, h=0.045)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — TITLE
# ═══════════════════════════════════════════════════════════════════════════════
s = slide()
bar(s, 0, 0, 13.33, color=CYAN, h=0.07)

# subtle ghost text
ghost = s.shapes.add_textbox(Inches(-0.5), Inches(1.2), Inches(14.5), Inches(5))
ghost.word_wrap = False
gf = ghost.text_frame
gp = gf.paragraphs[0]
gr = gp.add_run()
gr.text = "TRINETRA"
gr.font.size  = Pt(195)
gr.font.bold  = True
gr.font.color.rgb = RGBColor(0x11, 0x1A, 0x28)

# top-left accent block
rect(s, 0.5, 1.0, 0.09, 1.45, fill=CYAN)

label(s, "TRINETRA", 0.75, 0.92, 11, 1.0,
      size=52, bold=True, color=CYAN)
label(s, "CHAOS CONSOLE", 0.76, 1.82, 11, 0.75,
      size=32, bold=True, color=WHITE)

bar(s, 0.76, 2.65, 10.5, GOLD, h=0.06)

label(s, "Adaptive Consensus Under Network Failures & Malicious Validators",
      0.76, 2.8, 11, 0.55, size=17, italic=True, color=SILVER)

label(s, "A flight simulator for blockchain consensus —\ncrash your validators, inject Byzantine attacks, earn a certified Resilience Score.",
      0.76, 3.48, 10.5, 0.9, size=16, color=WHITE)

# 3 mini pills
tags = [("Real Ed25519 Crypto", CYAN), ("Full Raft Protocol", GOLD), ("Published Score Formula", GREEN)]
for i, (t, c) in enumerate(tags):
    pill(s, t, 0.76 + i*3.6, 4.65, 3.3, 0.42, bg=CARD2, fg=c, size=12)

label(s, "Hackathon  ·  Innovise 2026  ·  Blockchain Track PS3",
      0.76, 5.35, 9, 0.38, size=12, color=DIM)

rect(s, 0, 7.18, 13.33, 0.32, fill=CARD)
label(s, "TRINETRA CHAOS CONSOLE  ·  Innovise 2026  ·  Blockchain PS3",
      0.35, 7.21, 9, 0.26, size=8, color=DIM)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — PROBLEM + SOLUTION
# ═══════════════════════════════════════════════════════════════════════════════
s = slide()
section_header(s, "SLIDE 2 / 6", "Problem & Solution", RED)

# left: problem
rect(s, 0.35, 1.4, 5.9, 5.55, fill=CARD)
bar(s, 0.35, 1.4, 5.9, RED, h=0.055)
label(s, "⚡  The Problem", 0.55, 1.5, 5.5, 0.5, size=15, bold=True, color=RED)

problems = [
    ("💥", "Servers Crash",
     "Leader node dies mid-block. Does the cluster elect a new leader in ms — or freeze forever?"),
    ("🔌", "Network Splits",
     "Two node groups can't talk. Do they write conflicting ledgers? Split-brain = double-spend."),
    ("🤥", "Byzantine Nodes",
     "A corrupt validator sends different data to different peers — an equivocation attack."),
]
for i, (icon, h, body) in enumerate(problems):
    t = 2.1 + i * 1.58
    rect(s, 0.5, t, 5.55, 1.42, fill=BG2)
    bar(s, 0.5, t, 5.55, RED, h=0.04)
    label(s, icon, 0.62, t+0.1, 0.5, 0.5, size=20)
    label(s, h, 1.15, t+0.1, 4.7, 0.38, size=13, bold=True, color=WHITE)
    label(s, body, 1.15, t+0.5, 4.7, 0.82, size=11.5, color=SILVER)

label(s, "🏦  No safe way to test these scenarios before production deployment.",
      0.5, 6.55, 5.7, 0.35, size=11, italic=True, color=GOLD)

# right: solution
rect(s, 6.6, 1.4, 6.38, 5.55, fill=CARD)
bar(s, 6.6, 1.4, 6.38, CYAN, h=0.055)
label(s, "✈️  Our Solution", 6.8, 1.5, 6.0, 0.5, size=15, bold=True, color=CYAN)

label(s, "A flight simulator for distributed consensus.",
      6.8, 2.05, 6.0, 0.45, size=13, bold=True, color=WHITE)
label(s, "Trinetra runs a real 5-node Raft cluster with Ed25519-signed blocks, then lets you break it every possible way — safely in simulation.",
      6.8, 2.52, 6.0, 0.7, size=12, color=SILVER)

pillars = [
    ("🔬  SIMULATE", CYAN,
     "5 validators · 150 tx/sec · Merkle-chained blocks · real crypto"),
    ("💣  CHAOS", RED,
     "Crash · Partition · Equivocation · Gray Failure · Latency Storm"),
    ("📊  CERTIFY", GREEN,
     "Resilience Score A–F · MTTR < 400 ms · 17 CI-gated tests"),
]
for i, (title, color, body) in enumerate(pillars):
    t = 3.4 + i*1.27
    rect(s, 6.75, t, 6.08, 1.15, fill=BG2)
    bar(s, 6.75, t, 6.08, color, h=0.04)
    label(s, title, 6.9, t+0.08, 5.8, 0.35, size=12, bold=True, color=color)
    label(s, body,  6.9, t+0.46, 5.8, 0.58, size=11, color=SILVER)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — ARCHITECTURE
# ═══════════════════════════════════════════════════════════════════════════════
s = slide()
section_header(s, "SLIDE 3 / 6", "System Architecture", GOLD)

layers = [
    # (l, t, w, h, accent, title, body)
    (0.35, 1.4,  12.63, 0.78, CYAN,
     "🖥️  WEB DASHBOARD",
     "Canvas topology · Chart.js TPS/latency · Scenario Composer (drag-drop) · Tx Studio · Voice Narration · Dark/Light Theme"),
    (0.35, 2.3,  12.63, 0.78, GOLD,
     "⚡  FASTAPI  REST + WebSocket  (10 Hz push)",
     "/inject · /scenario · /manual-tx · /chaos-* · /metrics · /report  —  Pydantic models · Prometheus exporter"),
]
for (l, t, w, h, ac, ti, bo) in layers:
    rect(s, l, t, w, h, fill=CARD)
    bar(s, l, t, w, ac, h=0.055)
    label(s, ti, l+0.18, t+0.09, w-0.35, 0.35, size=12, bold=True, color=ac)
    label(s, bo, l+0.18, t+0.45, w-0.35, 0.28, size=10.5, color=SILVER)

# bottom two side-by-side engine blocks
engines = [
    (0.35, 3.2,  6.08, 2.0, GREEN,
     "⚙️  CONSENSUS ENGINE",
     [" Raft:  PreVote · Leader Transfer · Planned Handoff",
      " Adaptive Timeouts  (Jacobson/Karels SRTT/RTTVAR/RTO)",
      " Byzantine Gossip + 4-of-5 Echo Quorum (Hardened Mode)",
      " Gray Failure Detection via EWMA jitter tracking"]),
    (6.6,  3.2,  6.08, 2.0, RED,
     "💣  CHAOS ENGINE",
     [" Chaos Monkey — 1000 Jepsen-style schedules",
      " Geo-Latency Presets  (Mumbai · Frankfurt · Virginia)",
      " Partition · Equivocation · Latency Storm · Slow Leader",
      " Autopilot — 6-chapter choreographed demo controller"]),
]
for (l, t, w, h, ac, ti, pts) in engines:
    rect(s, l, t, w, h, fill=CARD)
    bar(s, l, t, w, ac, h=0.055)
    label(s, ti, l+0.18, t+0.09, w-0.3, 0.38, size=12, bold=True, color=ac)
    multiline(s, pts, l+0.18, t+0.52, w-0.28, h-0.65, size=11, color=SILVER, gap=5)

# bottom ledger + observability
ledger_items = [
    (0.35, 5.35, 6.08, 1.6, DIM,
     "📒  LEDGER  &  INVARIANTS",
     [" Signed Txs · Merkle Blocks · TokenLedger state machine",
      " Nonce replay defence · Overdraft protection",
      " Continuous invariant checker — fork detection every tick"]),
    (6.6,  5.35, 6.08, 1.6, PURPLE,
     "📡  OBSERVABILITY",
     [" Prometheus /metrics  ·  Grafana dashboard JSON",
      " Wireshark-style RPC tracer  ·  Resilience Score A–F",
      " HTML + JSON certification report  ·  GitHub Actions CI"]),
]
for (l, t, w, h, ac, ti, pts) in ledger_items:
    rect(s, l, t, w, h, fill=CARD)
    bar(s, l, t, w, ac, h=0.055)
    label(s, ti, l+0.18, t+0.09, w-0.3, 0.38, size=12, bold=True, color=ac)
    multiline(s, pts, l+0.18, t+0.52, w-0.28, h-0.65, size=11, color=SILVER, gap=5)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — TECHNICAL DEPTH + KEY FEATURES
# ═══════════════════════════════════════════════════════════════════════════════
s = slide()
section_header(s, "SLIDE 4 / 6", "Technical Depth & Features", PURPLE)

# left: tech pillars
tech = [
    ("🔐", "Ed25519 Digital Signatures", CYAN,
     "Every transaction and block carries a real PyNaCl signature. Forgery is computationally infeasible."),
    ("🗳️", "Full Raft — PreVote Protocol", GOLD,
     "Prevents disruptive elections. Leader Transfer (TimeoutNow) for planned handoff. Log-matching invariant checked continuously."),
    ("📐", "Jacobson/Karels RTT Estimator", GREEN,
     "SRTT · RTTVAR · RTO — same algorithm as TCP. Adaptive election timeouts track per-follower jitter."),
    ("🌳", "Binary Merkle Trees", PURPLE,
     "SHA-256 canonical hashing. Inclusion proofs verify any tx without reading the full block."),
    ("🛡️", "Byzantine Echo Quorum", RED,
     "4-of-5 signed-header protocol. EquivProof generation + peer blacklisting blocks equivocation forks."),
]
for i, (icon, title, color, body) in enumerate(tech):
    t = 1.45 + i * 1.12
    rect(s, 0.35, t, 6.08, 1.0, fill=CARD)
    bar(s, 0.35, t, 0.06, color, h=1.0)
    label(s, icon,  0.52, t+0.1,  0.55, 0.5, size=20)
    label(s, title, 1.1,  t+0.08, 5.1,  0.36, size=12, bold=True, color=color)
    label(s, body,  1.1,  t+0.46, 5.1,  0.44, size=10.5, color=SILVER)

# right: features grid (14 features in 2 micro columns)
feats = [
    ("🎙️ Live Narrator",       CYAN),
    ("🔊 Voice Narration",      CYAN),
    ("🤖 Autopilot Demo",       GOLD),
    ("⏸️ Pause/Step/Replay",    GOLD),
    ("🌍 Geo-Latency Presets",  GREEN),
    ("📊 Resilience Score A–F", GREEN),
    ("🎭 Scenario Composer",    PURPLE),
    ("💸 Tx Studio",            PURPLE),
    ("🕵️ Gray Failure Detect",  RED),
    ("🧬 Merkle Proof Viewer",  GOLD),
    ("📡 Wireshark Tracer",     CYAN),
    ("🦊 Chaos Monkey",         RED),
    ("📈 Prometheus+Grafana",   GREEN),
    ("🌗 Dark / Light Theme",   SILVER),
]
label(s, "14 Advanced Features", 6.6, 1.38, 6.5, 0.42,
      size=14, bold=True, color=WHITE)
bar(s, 6.6, 1.8, 6.5, PURPLE, h=0.04)

col_w = 3.1
per_col = 7
for i, (name, color) in enumerate(feats):
    col = i // per_col
    row = i % per_col
    l = 6.65 + col * (col_w + 0.2)
    t = 1.95 + row * 0.74
    rect(s, l, t, col_w, 0.62, fill=CARD2)
    bar(s, l, t, 0.05, color, h=0.62)
    label(s, name, l+0.18, t+0.12, col_w-0.25, 0.38, size=11, color=color)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — RESULTS & LIVE DEMO
# ═══════════════════════════════════════════════════════════════════════════════
s = slide()
section_header(s, "SLIDE 5 / 6", "Results & Live Demo Walk-Through", GREEN)

# KPI row
kpis = [
    ("87.5", "/ 100",   "Resilience Score",    GOLD),
    ("<400", "ms",      "MTTR (Leader Crash)",  GREEN),
    ("150+", "tx / s",  "Sustained Throughput", CYAN),
    ("17",   "tests",   "All Passing — CI Gate",GREEN),
    ("-99%", "churn",   "Adaptive vs Static",   PURPLE),
]
kw = 2.42
for i, (big, unit, sub, ac) in enumerate(kpis):
    kpi_box(s, big, unit, sub, 0.35 + i*(kw+0.12), 1.35, kw, 1.62, ac)

# Benchmark comparison
label(s, "Adaptive vs. Static Benchmark", 0.35, 3.15, 7, 0.38,
      size=13, bold=True, color=WHITE)

hcols = ["Metric", "Static Mode", "Adaptive (Ours)", "Delta"]
hws   = [3.3, 2.4, 3.1, 2.5]
hls   = [0.35, 3.65, 6.05, 9.15]
for j, (h, hw, hl) in enumerate(zip(hcols, hws, hls)):
    rect(s, hl, 3.55, hw-0.05, 0.38,
         fill=CYAN if j==2 else CARD)
    label(s, h, hl+0.1, 3.56, hw-0.15, 0.36,
          size=11, bold=True,
          color=BG if j==2 else CYAN)

rows = [
    ["Resilience Score",  "57.8 / 100", "87.5 / 100", "+51% ✅"],
    ["Election Churn",    "618 / min",   "3 / min",    "−99.5% ✅"],
    ["p95 Latency",       "48 ms",       "48 ms",       "Stable ✅"],
    ["CI Gate  (≥ 80)",   "FAIL  ✗",     "PASS  ✓",     ""],
]
for ri, row in enumerate(rows):
    bg = CARD2 if ri % 2 == 0 else CARD
    for j, (cell, hw, hl) in enumerate(zip(row, hws, hls)):
        rect(s, hl, 3.95+ri*0.44, hw-0.05, 0.42, fill=bg)
        clr = GREEN if "✅" in cell or "✓" in cell else \
              RED   if "✗" in cell else WHITE
        label(s, cell, hl+0.1, 3.97+ri*0.44, hw-0.15, 0.38, size=11, color=clr)

# Demo steps right column
label(s, "Demo Walk-Through", 9.5, 3.15, 3.7, 0.38,
      size=13, bold=True, color=WHITE)
bar(s, 9.5, 3.55, 3.7, GREEN, h=0.04)

steps = [
    (GREEN,  "①  Cluster boots — 150 tx/s, crown on N0"),
    (RED,    "②  Click 'Crash N0' — heartbeats stop"),
    (GOLD,   "③  N2 wins election in ~320 ms"),
    (CYAN,   "④  Recover All — N0 rejoins as Follower"),
    (PURPLE, "⑤  Equivocation blocked by Echo Quorum"),
    (GOLD,   "⑥  Autopilot → Score: A  87.5 / 100"),
]
for i, (c, txt_) in enumerate(steps):
    t = 3.62 + i * 0.57
    rect(s, 9.5, t, 3.7, 0.5, fill=CARD2)
    bar(s, 9.5, t, 0.05, c, h=0.5)
    label(s, txt_, 9.65, t+0.08, 3.45, 0.34, size=10.5, color=WHITE)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — WHY TOP 1  /  CLOSING
# ═══════════════════════════════════════════════════════════════════════════════
s = slide()
bar(s, 0, 0, 13.33, color=GOLD, h=0.07)
rect(s, 0, 7.18, 13.33, 0.32, fill=CARD)
label(s, "TRINETRA CHAOS CONSOLE  ·  Innovise 2026  ·  Blockchain PS3",
      0.35, 7.21, 9, 0.26, size=8, color=DIM)

# ghost
ghost2 = s.shapes.add_textbox(Inches(-0.3), Inches(1.8), Inches(14), Inches(5))
ghost2.word_wrap = False
gf2 = ghost2.text_frame
gp2 = gf2.paragraphs[0]
gr2 = gp2.add_run()
gr2.text = "TOP 1"
gr2.font.size  = Pt(210)
gr2.font.bold  = True
gr2.font.color.rgb = RGBColor(0x11, 0x1A, 0x27)

label(s, "SLIDE 6 / 6", 0.38, 0.12, 3, 0.3, size=9, bold=True, color=GOLD)
label(s, "Why Trinetra Wins", 0.38, 0.42, 12, 0.65, size=28, bold=True, color=WHITE)
bar(s, 0.38, 1.1, 12.6, GOLD, h=0.055)

# vs table
label(s, "Most Hackathon Projects", 0.5, 1.28, 5.9, 0.42, size=14, bold=True, color=RED)
label(s, "Trinetra Chaos Console", 7.0, 1.28, 5.9, 0.42, size=14, bold=True, color=GREEN)

vs_rows = [
    ("Fake random numbers for 'consensus'",      "Real Raft — PreVote · Leader Transfer · Gray Failure"),
    ("No cryptography (hashes only)",            "Ed25519 on every tx & block header (PyNaCl)"),
    ("Single happy-path demo",                   "8 named scenarios + drag-and-drop Scenario Composer"),
    ("No automated tests",                       "17 pytest tests + GitHub Actions CI gate"),
    ("Black-box score, no formula",              "Published formula — Score = f(Churn, MTTR, Avail, Forks)"),
    ("No observability",                         "Prometheus + Grafana + Wireshark-style RPC tracer"),
]
for i, (bad, good) in enumerate(vs_rows):
    t = 1.78 + i * 0.67
    bg = CARD2 if i % 2 == 0 else CARD
    rect(s, 0.38, t, 5.9, 0.58, fill=bg)
    rect(s, 6.95, t, 5.9, 0.58, fill=bg)
    label(s, "✗  " + bad,  0.52, t+0.1, 5.6, 0.42, size=11.5, color=RGBColor(0xFF,0x88,0x88))
    label(s, "✅  " + good, 7.08, t+0.1, 5.65, 0.42, size=11.5, color=WHITE)

# VS divider
rect(s, 6.55, 1.25, 0.04, 4.55, fill=DIM)
label(s, "VS", 6.28, 3.2, 0.7, 0.45, size=18, bold=True,
      color=DIM, align=PP_ALIGN.CENTER)

# bottom CTA bar
rect(s, 0.35, 6.05, 12.65, 0.95, fill=CARD)
bar(s, 0.35, 6.05, 12.65, GOLD, h=0.055)
label(s, "🏆  Judges can verify every claim live on screen — no slides, just working software.",
      0.55, 6.12, 8.5, 0.38, size=13, bold=True, color=GOLD)
label(s, "🔗  github.com/jaydevgania/Trinetra",
      0.55, 6.5, 6, 0.35, size=12, color=CYAN)
label(s, "🌐  http://127.0.0.1:8000  (live demo)",
      6.0, 6.5, 6.5, 0.35, size=12, color=CYAN, align=PP_ALIGN.RIGHT)


# ── Save ──────────────────────────────────────────────────────────────────────
out = r"D:\VJTI\Trinetra_Submission.pptx"
prs.save(out)
print(f"Saved  →  {out}")
print(f"Slides →  {len(prs.slides)}")
